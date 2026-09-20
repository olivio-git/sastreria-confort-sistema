"""
test_pago_comision_seleccion.py
================================
Cubre el pago de comisión por selección de devengaciones puntuales
(AplicacionPagoComision) introducido junto con la mejora de la sección
Comisiones del detalle de empleado:

  - comisiones.estado_asignacion: pagada / parcial / pendiente.
  - pagar_comision_empleado: paga exactamente lo seleccionado, crea
    PagoComisionEmpleado + AplicacionPagoComision + CajaMovimiento (si
    via_caja), todo en una transacción.
  - Rechazo de selección: clave de otro empleado, ya pagada, clave
    manipulada/con formato inválido, selección vacía.
  - Pago parcial: sólo cubre el remanente, no el monto_comision_fijo completo.
  - comisiones.asignar_pagos_fifo: el algoritmo de reparto FIFO que también
    usa la migración de datos 0066, probado con modelos reales.
  - Excel: `?estado=pendientes` filtra sólo pendientes/parciales,
    `?estado=pagadas` sólo pagadas; color de fila por estado real, hoja
    "Pagadas" fija y línea de reconciliación (Pendiente bruto − Saldo a
    favor = Saldo pendiente).
  - detalle_empleado: 200 con el filtro "Pendientes" activo por defecto;
    conteos por estado real (pendiente/parcial/pagada) y tinte de fila.
"""
from decimal import Decimal
from io import BytesIO
from unittest import mock
import importlib

import openpyxl
from django.apps import apps as django_apps
from django.contrib.auth.models import User
from django.contrib.messages import get_messages
from django.db import OperationalError
from django.db.models import Sum
from django.test import Client, TestCase
from django.urls import reverse

from misastreria import comisiones
from misastreria.models import (
    AplicacionPagoComision, CajaMovimiento, PagoComisionEmpleado,
    ReparacionEmpleado,
)
from misastreria.views import ASIGNACION_MODELOS, _devengaciones_empleado

from .factories import make_cliente, make_empleado, make_reparacion, make_sesion_caja, make_user


def _crear_asignacion_reparacion(empleado, monto, **kwargs):
    rep = make_reparacion(**kwargs)
    return ReparacionEmpleado.objects.create(
        reparacion=rep, empleado=empleado, monto_comision_fijo=monto)


# ──────────────────────────────────────────────────────────────────────────────
# Estado de una devengación
# ──────────────────────────────────────────────────────────────────────────────

class EstadoAsignacionTests(TestCase):

    def test_pendiente_sin_pago(self):
        self.assertEqual(comisiones.estado_asignacion(Decimal('100'), Decimal('0')), 'pendiente')

    def test_parcial(self):
        self.assertEqual(comisiones.estado_asignacion(Decimal('100'), Decimal('40')), 'parcial')

    def test_pagada_exacto(self):
        self.assertEqual(comisiones.estado_asignacion(Decimal('100'), Decimal('100')), 'pagada')

    def test_pagada_si_pagado_excede(self):
        # No debería pasar en la práctica (nunca se aplica más que el pendiente),
        # pero no debe reventar si pasa: se considera pagada.
        self.assertEqual(comisiones.estado_asignacion(Decimal('100'), Decimal('120')), 'pagada')

    def test_estado_filtro_agrupa_parcial_con_pendiente(self):
        self.assertEqual(comisiones.estado_filtro('pendiente'), 'pendiente')
        self.assertEqual(comisiones.estado_filtro('parcial'), 'pendiente')
        self.assertEqual(comisiones.estado_filtro('pagada'), 'pagada')


# ──────────────────────────────────────────────────────────────────────────────
# pagar_comision_empleado — pago por selección
# ──────────────────────────────────────────────────────────────────────────────

class PagarComisionSeleccionTests(TestCase):

    def setUp(self):
        self.user = make_user()
        self.client = Client()
        self.client.force_login(self.user)
        self.empleado = make_empleado()

    def _url(self, empleado=None):
        return reverse('pagar_comision_empleado', kwargs={'empleado_id': (empleado or self.empleado).id})

    def test_pagar_una_devengacion_completa(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('150'))
        resp = self.client.post(self._url(), {
            'sel': [f'reparacion:{a.id}'],
            'forma_pago': 'efectivo',
            'via_caja': '',
            'descripcion': 'Pago de prueba',
        })
        self.assertEqual(resp.status_code, 302)

        pago = PagoComisionEmpleado.objects.get(empleado=self.empleado)
        self.assertEqual(pago.monto, Decimal('150'))
        aplicaciones = list(pago.aplicaciones.all())
        self.assertEqual(len(aplicaciones), 1)
        self.assertEqual(aplicaciones[0].reparacion_empleado_id, a.id)
        self.assertEqual(aplicaciones[0].monto, Decimal('150'))
        # Sin via_caja: no debe crear movimiento de caja.
        self.assertFalse(CajaMovimiento.objects.filter(referencia_pago_comision=pago).exists())

    def test_pagar_varias_devengaciones_suma_el_monto(self):
        a1 = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        a2 = _crear_asignacion_reparacion(self.empleado, Decimal('50'))
        resp = self.client.post(self._url(), {
            'sel': [f'reparacion:{a1.id}', f'reparacion:{a2.id}'],
            'forma_pago': 'efectivo',
            'via_caja': '',
        })
        self.assertEqual(resp.status_code, 302)
        pago = PagoComisionEmpleado.objects.get(empleado=self.empleado)
        self.assertEqual(pago.monto, Decimal('150'))
        self.assertEqual(pago.aplicaciones.count(), 2)

    def test_pagar_via_caja_crea_movimiento(self):
        make_sesion_caja(usuario=self.user)
        a = _crear_asignacion_reparacion(self.empleado, Decimal('80'))
        resp = self.client.post(self._url(), {
            'sel': [f'reparacion:{a.id}'],
            'forma_pago': 'efectivo',
            'via_caja': 'on',
        })
        self.assertEqual(resp.status_code, 302)
        pago = PagoComisionEmpleado.objects.get(empleado=self.empleado)
        mov = CajaMovimiento.objects.get(referencia_pago_comision=pago)
        self.assertEqual(mov.monto, Decimal('80'))
        self.assertEqual(mov.concepto, 'comision_empleado')

    def test_pagar_parcial_solo_cubre_el_remanente(self):
        """Una devengación con Bs 100 ya pagados Bs 60: seleccionarla de nuevo
        sólo debe cobrar los Bs 40 pendientes, no los 100 completos."""
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        pago_previo = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('60'))
        AplicacionPagoComision.objects.create(pago=pago_previo, reparacion_empleado=a, monto=Decimal('60'))

        resp = self.client.post(self._url(), {
            'sel': [f'reparacion:{a.id}'],
            'forma_pago': 'efectivo',
            'via_caja': '',
        })
        self.assertEqual(resp.status_code, 302)
        nuevo_pago = PagoComisionEmpleado.objects.exclude(pk=pago_previo.pk).get(empleado=self.empleado)
        self.assertEqual(nuevo_pago.monto, Decimal('40'))
        ap = nuevo_pago.aplicaciones.get()
        self.assertEqual(ap.monto, Decimal('40'))

    def test_rechaza_devengacion_de_otro_empleado(self):
        otro = make_empleado(ci='OTRO-1')
        a = _crear_asignacion_reparacion(otro, Decimal('100'))
        resp = self.client.post(self._url(self.empleado), {
            'sel': [f'reparacion:{a.id}'],
            'forma_pago': 'efectivo',
            'via_caja': '',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(PagoComisionEmpleado.objects.filter(empleado=self.empleado).exists())
        # La devengación del OTRO empleado tampoco debe haber sido tocada.
        self.assertFalse(AplicacionPagoComision.objects.filter(reparacion_empleado=a).exists())

    def test_rechaza_devengacion_ya_pagada(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        pago_previo = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('100'))
        AplicacionPagoComision.objects.create(pago=pago_previo, reparacion_empleado=a, monto=Decimal('100'))

        resp = self.client.post(self._url(), {
            'sel': [f'reparacion:{a.id}'],
            'forma_pago': 'efectivo',
            'via_caja': '',
        })
        self.assertEqual(resp.status_code, 302)
        # No se creó un segundo pago.
        self.assertEqual(PagoComisionEmpleado.objects.filter(empleado=self.empleado).count(), 1)

    def test_rechaza_clave_con_formato_invalido(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        for clave_mala in ['reparacion', 'reparacion:abc', 'inexistente:1', f'reparacion:{a.id}:extra']:
            resp = self.client.post(self._url(), {
                'sel': [clave_mala],
                'forma_pago': 'efectivo',
                'via_caja': '',
            })
            self.assertEqual(resp.status_code, 302)
        self.assertFalse(PagoComisionEmpleado.objects.filter(empleado=self.empleado).exists())

    def test_rechaza_seleccion_vacia(self):
        resp = self.client.post(self._url(), {
            'sel': [],
            'forma_pago': 'efectivo',
            'via_caja': '',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(PagoComisionEmpleado.objects.filter(empleado=self.empleado).exists())

    def test_una_clave_invalida_rechaza_toda_la_seleccion(self):
        """Si UNA clave de la selección es inválida, no se paga NADA — ni las
        demás claves válidas del mismo POST."""
        valida = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        otro = make_empleado(ci='OTRO-2')
        invalida = _crear_asignacion_reparacion(otro, Decimal('50'))

        resp = self.client.post(self._url(), {
            'sel': [f'reparacion:{valida.id}', f'reparacion:{invalida.id}'],
            'forma_pago': 'efectivo',
            'via_caja': '',
        })
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(PagoComisionEmpleado.objects.filter(empleado=self.empleado).exists())
        self.assertFalse(AplicacionPagoComision.objects.filter(reparacion_empleado=valida).exists())

    def test_get_no_permitido(self):
        resp = self.client.get(self._url())
        self.assertEqual(resp.status_code, 405)


# ──────────────────────────────────────────────────────────────────────────────
# comisiones.asignar_pagos_fifo — algoritmo puro (también usado por la 0066)
# ──────────────────────────────────────────────────────────────────────────────

class _PagoFake:
    """Objeto mínimo con `.monto`, para no depender de guardar en DB en cada
    escenario de reparto FIFO."""
    def __init__(self, monto):
        self.monto = Decimal(monto)


class AsignarPagosFifoTests(TestCase):
    """Prueba el algoritmo con modelos reales (AplicacionPagoComision se
    persiste de verdad vía el callable `aplicar`), como pide el criterio de
    aceptación: la migración 0066 y estos tests comparten la misma función."""

    def setUp(self):
        self.empleado = make_empleado()

    def _accrual(self, monto, saldo=None):
        a = _crear_asignacion_reparacion(self.empleado, monto)
        creadas = []

        def aplicar(pago, cantidad):
            creadas.append(AplicacionPagoComision.objects.create(
                pago=pago, reparacion_empleado=a, monto=cantidad))

        return {'saldo': Decimal(saldo if saldo is not None else monto), 'aplicar': aplicar}, a, creadas

    def test_un_pago_cubre_una_devengacion_exacta(self):
        acc, asignacion, creadas = self._accrual(Decimal('100'))
        pago = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('100'))
        sobrante = comisiones.asignar_pagos_fifo([pago], [acc])
        self.assertEqual(sobrante, Decimal('0'))
        self.assertEqual(len(creadas), 1)
        self.assertEqual(creadas[0].monto, Decimal('100'))
        self.assertEqual(acc['saldo'], Decimal('0'))

    def test_un_pago_se_reparte_entre_dos_devengaciones_oldest_first(self):
        acc1, _, creadas1 = self._accrual(Decimal('60'))
        acc2, _, creadas2 = self._accrual(Decimal('60'))
        pago = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('100'))
        sobrante = comisiones.asignar_pagos_fifo([pago], [acc1, acc2])
        self.assertEqual(sobrante, Decimal('0'))
        self.assertEqual(creadas1[0].monto, Decimal('60'))  # la más vieja se cubre entera
        self.assertEqual(creadas2[0].monto, Decimal('40'))  # la siguiente recibe el resto
        self.assertEqual(acc1['saldo'], Decimal('0'))
        self.assertEqual(acc2['saldo'], Decimal('20'))

    def test_varios_pagos_no_retroceden_sobre_devengacion_ya_cubierta(self):
        acc1, _, creadas1 = self._accrual(Decimal('50'))
        acc2, _, creadas2 = self._accrual(Decimal('50'))
        pago1 = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('50'))
        pago2 = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('50'))
        comisiones.asignar_pagos_fifo([pago1, pago2], [acc1, acc2])
        self.assertEqual(sum(a.monto for a in creadas1), Decimal('50'))
        self.assertEqual(sum(a.monto for a in creadas2), Decimal('50'))

    def test_sobrante_cuando_los_pagos_exceden_lo_devengado(self):
        acc, _, creadas = self._accrual(Decimal('30'))
        pago = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('50'))
        sobrante = comisiones.asignar_pagos_fifo([pago], [acc])
        self.assertEqual(sobrante, Decimal('20'))
        self.assertEqual(creadas[0].monto, Decimal('30'))

    def test_saldo_de_arranque_ya_reducido_es_idempotente(self):
        """Simula una corrida previa de la migración: el saldo de arranque ya
        viene descontado de lo aplicado antes, así que un segundo pago no
        debe volver a cubrir lo que el primero ya cubrió."""
        acc, asignacion, creadas = self._accrual(Decimal('100'), saldo=Decimal('40'))
        pago_nuevo = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('40'))
        comisiones.asignar_pagos_fifo([pago_nuevo], [acc])
        self.assertEqual(len(creadas), 1)
        self.assertEqual(creadas[0].monto, Decimal('40'))
        self.assertEqual(acc['saldo'], Decimal('0'))


# ──────────────────────────────────────────────────────────────────────────────
# Excel — filtro de pendientes
# ──────────────────────────────────────────────────────────────────────────────

class ExcelPendientesTests(TestCase):

    def setUp(self):
        self.user = make_user()
        self.client = Client()
        self.client.force_login(self.user)
        self.empleado = make_empleado()

    def test_excel_pendientes_filtra_filename(self):
        _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        url = reverse('exportar_devengaciones_empleado_excel', kwargs={'id': self.empleado.id})
        resp = self.client.get(url, {'estado': 'pendientes'})
        self.assertEqual(resp.status_code, 200)
        self.assertIn('devengaciones_pendientes_', resp['Content-Disposition'])

    def test_excel_sin_filtro_no_lleva_pendientes_en_el_nombre(self):
        _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        url = reverse('exportar_devengaciones_empleado_excel', kwargs={'id': self.empleado.id})
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertNotIn('pendientes', resp['Content-Disposition'])


# ──────────────────────────────────────────────────────────────────────────────
# Excel — colores por estado, hoja "Pagadas", filtro `?estado=pagadas` y
# reconciliación (Pendiente bruto − Saldo a favor = Saldo pendiente)
# ──────────────────────────────────────────────────────────────────────────────

class ExcelColoresYPagadasTests(TestCase):

    def setUp(self):
        self.user = make_user()
        self.client = Client()
        self.client.force_login(self.user)
        self.empleado = make_empleado()
        # Una fila de cada estado real: pendiente, parcial y pagada.
        self.a_pendiente = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        self.a_parcial = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        pago_parcial = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('40'))
        AplicacionPagoComision.objects.create(
            pago=pago_parcial, monto=Decimal('40'), reparacion_empleado=self.a_parcial)
        self.a_pagada = _crear_asignacion_reparacion(self.empleado, Decimal('60'))
        pago_pagada = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('60'))
        AplicacionPagoComision.objects.create(
            pago=pago_pagada, monto=Decimal('60'), reparacion_empleado=self.a_pagada)

    def _url(self, **params):
        url = reverse('exportar_devengaciones_empleado_excel', kwargs={'id': self.empleado.id})
        if params:
            qs = '&'.join(f'{k}={v}' for k, v in params.items())
            return f'{url}?{qs}'
        return url

    def test_filas_de_devengaciones_llevan_el_color_de_su_estado(self):
        resp = self.client.get(self._url())
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb['Devengaciones']
        colores_por_codigo = {}
        for row in ws.iter_rows(min_row=1):
            for celda in row:
                if celda.value in (self.a_pendiente.reparacion.codigo, self.a_parcial.reparacion.codigo,
                                    self.a_pagada.reparacion.codigo):
                    fila_celdas = ws[celda.row]
                    colores_por_codigo[celda.value] = fila_celdas[0].fill.fgColor.rgb
        self.assertTrue(colores_por_codigo[self.a_pendiente.reparacion.codigo].endswith('FEF3C7'))
        self.assertTrue(colores_por_codigo[self.a_parcial.reparacion.codigo].endswith('DBEAFE'))
        self.assertTrue(colores_por_codigo[self.a_pagada.reparacion.codigo].endswith('D1FAE5'))

    def test_estado_pagadas_filtra_solo_filas_pagadas(self):
        resp = self.client.get(self._url(estado='pagadas'))
        self.assertEqual(resp.status_code, 200)
        self.assertIn('devengaciones_pagadas_', resp['Content-Disposition'])
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb['Devengaciones']
        codigos_en_hoja = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertTrue(any(self.a_pagada.reparacion.codigo in v for v in codigos_en_hoja))
        self.assertFalse(any(self.a_pendiente.reparacion.codigo in v for v in codigos_en_hoja))
        self.assertFalse(any(self.a_parcial.reparacion.codigo in v for v in codigos_en_hoja))

    def test_hoja_pagadas_existe_y_lista_solo_lo_pagado(self):
        resp = self.client.get(self._url())
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        self.assertIn('Pagadas', wb.sheetnames)
        ws = wb['Pagadas']
        codigos_en_hoja = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertTrue(any(self.a_pagada.reparacion.codigo in v for v in codigos_en_hoja))
        self.assertFalse(any(self.a_pendiente.reparacion.codigo in v for v in codigos_en_hoja))
        self.assertFalse(any(self.a_parcial.reparacion.codigo in v for v in codigos_en_hoja))

    def test_hoja_pagadas_existe_incluso_descargando_pendientes(self):
        """La hoja "Pagadas" es una referencia fija: no depende del filtro
        `?estado=` que se esté aplicando a la hoja "Devengaciones"."""
        resp = self.client.get(self._url(estado='pendientes'))
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        self.assertIn('Pagadas', wb.sheetnames)
        ws = wb['Pagadas']
        codigos_en_hoja = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertTrue(any(self.a_pagada.reparacion.codigo in v for v in codigos_en_hoja))

    def test_reconciliacion_pendiente_bruto_menos_saldo_a_favor(self):
        datos = _devengaciones_empleado(self.empleado)
        resp = self.client.get(self._url())
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        ws = wb['Devengaciones']
        valores = {c.value: ws.cell(row=c.row, column=2).value
                   for row in ws.iter_rows(min_row=1, max_col=1) for c in row if c.value}
        self.assertAlmostEqual(valores['Pendiente bruto'], float(datos['total_pendiente_bruto']))
        self.assertAlmostEqual(valores['Saldo a favor (no aplicado)'], float(datos['no_aplicado_comision']))
        self.assertAlmostEqual(valores['Saldo pendiente'], float(datos['saldo_comision']))
        # La identidad matemática que motivó el pedido del reviewer.
        self.assertAlmostEqual(
            valores['Pendiente bruto'] - valores['Saldo a favor (no aplicado)'],
            valores['Saldo pendiente'],
        )


# ──────────────────────────────────────────────────────────────────────────────
# Detalle de empleado — filtro Pendientes activo por defecto
# ──────────────────────────────────────────────────────────────────────────────

class DetalleEmpleadoComisionesTests(TestCase):

    def setUp(self):
        self.user = make_user()
        self.client = Client()
        self.client.force_login(self.user)
        self.empleado = make_empleado()

    def test_detalle_200_con_pendientes_activo_por_defecto(self):
        _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        url = reverse('detalle_empleado', kwargs={'id': self.empleado.id})
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode()
        self.assertIn('data-value="pendiente parcial"', html)
        self.assertIn('id="modalPagarComision"', html)

    def test_asignacion_modelos_cubre_los_5_tipos(self):
        self.assertEqual(
            set(ASIGNACION_MODELOS.keys()),
            {'reparacion', 'confeccion', 'venta', 'alquiler', 'produccion'},
        )

    def test_detalle_muestra_conteos_por_estado_real(self):
        _crear_asignacion_reparacion(self.empleado, Decimal('100'))  # pendiente
        parcial = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        pago_parcial = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('40'))
        AplicacionPagoComision.objects.create(
            pago=pago_parcial, monto=Decimal('40'), reparacion_empleado=parcial)
        pagada = _crear_asignacion_reparacion(self.empleado, Decimal('60'))
        pago_pagada = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('60'))
        AplicacionPagoComision.objects.create(
            pago=pago_pagada, monto=Decimal('60'), reparacion_empleado=pagada)

        datos = _devengaciones_empleado(self.empleado)
        self.assertEqual(datos['n_pendientes_comision'], 1)
        self.assertEqual(datos['n_parciales_comision'], 1)
        self.assertEqual(datos['n_pagadas_comision'], 1)
        self.assertEqual(datos['total_operaciones_comision'], 3)

        resp = self.client.get(reverse('detalle_empleado', kwargs={'id': self.empleado.id}))
        html = resp.content.decode()
        self.assertIn('data-value="parcial"', html)
        self.assertIn('data-value="pagada"', html)
        self.assertIn('sc-fila-pendiente', html)
        self.assertIn('sc-fila-parcial', html)
        self.assertIn('sc-fila-pagada', html)


# ──────────────────────────────────────────────────────────────────────────────
# Anti doble pago: claves duplicadas, saldo a favor, defensa por saldo
# ──────────────────────────────────────────────────────────────────────────────

class AntiDoblePagoTests(TestCase):

    def setUp(self):
        self.user = make_user()
        self.client = Client()
        self.client.force_login(self.user)
        self.empleado = make_empleado()

    def _url(self):
        return reverse('pagar_comision_empleado', kwargs={'empleado_id': self.empleado.id})

    def _pagar(self, *claves):
        return self.client.post(self._url(), {
            'sel': list(claves), 'forma_pago': 'efectivo', 'via_caja': '',
        })

    def _pago_con_huerfana(self, monto):
        pago = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal(monto))
        ap = AplicacionPagoComision.objects.create(
            pago=pago, monto=Decimal(monto), detalle_snapshot='Reparación borrada')
        return pago, ap

    # C1 ────────────────────────────────────────────────────────────────────
    def test_claves_duplicadas_y_con_ceros_se_pagan_una_sola_vez(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        resp = self._pagar(f'reparacion:{a.id}', f'reparacion:{a.id}', f'reparacion:0{a.id}', f'reparacion:00{a.id}')
        self.assertEqual(resp.status_code, 302)
        pago = PagoComisionEmpleado.objects.get(empleado=self.empleado)
        self.assertEqual(pago.monto, Decimal('100'))
        self.assertEqual(pago.aplicaciones.count(), 1)

    def test_parsear_clave_normaliza_y_rechaza(self):
        tipos = {'reparacion': None}
        self.assertEqual(comisiones.parsear_clave_devengacion('reparacion:05', tipos), ('reparacion', 5))
        for mala in ['reparacion:', 'reparacion:0', 'reparacion: 5', 'reparacion:٥', 'reparacion:-1', 'x:1', '', None]:
            with self.assertRaises(ValueError):
                comisiones.parsear_clave_devengacion(mala, tipos)

    def test_doble_submit_secuencial_no_paga_dos_veces(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        self._pagar(f'reparacion:{a.id}')
        self._pagar(f'reparacion:{a.id}')
        self.assertEqual(PagoComisionEmpleado.objects.filter(empleado=self.empleado).count(), 1)
        self.assertEqual(
            AplicacionPagoComision.objects.filter(reparacion_empleado=a).aggregate(
                s=Sum('monto'))['s'], Decimal('100'))

    # C2a ───────────────────────────────────────────────────────────────────
    def test_saldo_a_favor_cubre_todo_sin_pago_nuevo(self):
        pago_viejo, ap = self._pago_con_huerfana('100')
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        resp = self._pagar(f'reparacion:{a.id}')
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(PagoComisionEmpleado.objects.filter(empleado=self.empleado).count(), 1)
        ap.refresh_from_db()
        self.assertEqual(ap.reparacion_empleado_id, a.id)
        self.assertEqual(ap.pago_id, pago_viejo.id)
        self.assertFalse(CajaMovimiento.objects.exists())
        msgs = [str(m) for m in get_messages(resp.wsgi_request)]
        self.assertTrue(any('saldo a favor' in m and '100.00' in m for m in msgs), msgs)

    def test_saldo_a_favor_parcial_se_parte_la_aplicacion(self):
        pago_viejo, ap = self._pago_con_huerfana('100')
        a1 = _crear_asignacion_reparacion(self.empleado, Decimal('60'))
        self._pagar(f'reparacion:{a1.id}')
        # No hay pago nuevo: 60 de la huérfana se movieron a a1, quedan 40 a favor.
        self.assertEqual(PagoComisionEmpleado.objects.filter(empleado=self.empleado).count(), 1)
        ap.refresh_from_db()
        self.assertEqual(ap.monto, Decimal('40'))
        self.assertIsNone(ap.asignacion)
        nueva = AplicacionPagoComision.objects.get(reparacion_empleado=a1)
        self.assertEqual(nueva.monto, Decimal('60'))
        self.assertEqual(nueva.pago_id, pago_viejo.id)
        self.assertEqual(_devengaciones_empleado(self.empleado)['no_aplicado_comision'], Decimal('40'))

        # Segundo pago: consume los 40 restantes y cobra sólo 60 nuevos.
        a2 = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        resp = self._pagar(f'reparacion:{a2.id}')
        nuevo = PagoComisionEmpleado.objects.filter(empleado=self.empleado).order_by('-id').first()
        self.assertNotEqual(nuevo.pk, pago_viejo.pk)
        self.assertEqual(nuevo.monto, Decimal('60'))
        self.assertEqual(
            AplicacionPagoComision.objects.filter(reparacion_empleado=a2).aggregate(
                s=Sum('monto'))['s'], Decimal('100'))
        datos = _devengaciones_empleado(self.empleado)
        self.assertEqual(datos['no_aplicado_comision'], Decimal('0'))
        self.assertEqual(datos['saldo_comision'], Decimal('0'))
        msgs = [str(m) for m in get_messages(resp.wsgi_request)]
        self.assertTrue(any('40.00 cubiertos con saldo a favor' in m for m in msgs), msgs)

    def test_excedente_sobre_comision_reducida_se_mueve(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        self._pagar(f'reparacion:{a.id}')
        ReparacionEmpleado.objects.filter(pk=a.pk).update(monto_comision_fijo=Decimal('60'))
        b = _crear_asignacion_reparacion(self.empleado, Decimal('50'))
        self._pagar(f'reparacion:{b.id}')

        pagos = PagoComisionEmpleado.objects.filter(empleado=self.empleado).order_by('id')
        self.assertEqual([p.monto for p in pagos], [Decimal('100'), Decimal('10')])
        suma = lambda **kw: AplicacionPagoComision.objects.filter(**kw).aggregate(s=Sum('monto'))['s']
        self.assertEqual(suma(reparacion_empleado=a), Decimal('60'))
        self.assertEqual(suma(reparacion_empleado=b), Decimal('50'))
        datos = _devengaciones_empleado(self.empleado)
        self.assertEqual(datos['no_aplicado_comision'], Decimal('0'))
        self.assertEqual(datos['saldo_comision'], Decimal('0'))

    def test_remanente_de_pago_sin_aplicar_se_consume(self):
        legado = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('30'))
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        self.assertEqual(_devengaciones_empleado(self.empleado)['no_aplicado_comision'], Decimal('30'))
        self._pagar(f'reparacion:{a.id}')
        nuevo = PagoComisionEmpleado.objects.exclude(pk=legado.pk).get(empleado=self.empleado)
        self.assertEqual(nuevo.monto, Decimal('70'))
        self.assertEqual(legado.aplicaciones.get().monto, Decimal('30'))

    # C2b ───────────────────────────────────────────────────────────────────
    def test_defensa_rechaza_si_el_monto_nuevo_supera_el_saldo(self):
        """Si por lo que sea el saldo a favor no se consumiera, la defensa por
        saldo (devengado − pagado) debe impedir el pago."""
        self._pago_con_huerfana('100')
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        with mock.patch('misastreria.views._fuentes_saldo_a_favor', return_value=(Decimal('0'), [])):
            resp = self._pagar(f'reparacion:{a.id}')
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(PagoComisionEmpleado.objects.filter(empleado=self.empleado).count(), 1)
        self.assertFalse(AplicacionPagoComision.objects.filter(reparacion_empleado=a).exists())
        msgs = [str(m) for m in get_messages(resp.wsgi_request)]
        self.assertTrue(any('supera' in m for m in msgs), msgs)

    # W4 ────────────────────────────────────────────────────────────────────
    def test_deadlock_da_mensaje_y_no_500(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        with mock.patch('misastreria.views._resolver_clave_devengacion',
                        side_effect=OperationalError('Deadlock found')):
            resp = self._pagar(f'reparacion:{a.id}')
        self.assertEqual(resp.status_code, 302)
        self.assertFalse(PagoComisionEmpleado.objects.exists())
        msgs = [str(m) for m in get_messages(resp.wsgi_request)]
        self.assertTrue(any('Intenta de nuevo' in m for m in msgs), msgs)


# ──────────────────────────────────────────────────────────────────────────────
# Snapshot, template sin localizar, Excel seguro, FIFO por remanente, 0066
# ──────────────────────────────────────────────────────────────────────────────

class SnapshotYPresentacionTests(TestCase):

    def setUp(self):
        self.user = make_user()
        self.client = Client()
        self.client.force_login(self.user)
        self.empleado = make_empleado()

    def test_snapshot_describe_la_operacion_y_se_trunca(self):
        cli = make_cliente(nombres='X' * 150, apellido_paterno='Y' * 90)
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'), cliente=cli)
        texto = comisiones.detalle_snapshot('reparacion', a, max_length=200)
        self.assertLessEqual(len(texto), 200)
        self.assertTrue(texto.startswith(f'Reparación {a.reparacion.codigo}'))
        corto = comisiones.detalle_snapshot('reparacion', _crear_asignacion_reparacion(self.empleado, Decimal('5')))
        self.assertIn('comisión Bs 5.00', corto)

    def test_pago_guarda_snapshot_legible(self):
        a = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        self.client.post(reverse('pagar_comision_empleado', kwargs={'empleado_id': self.empleado.id}),
                         {'sel': [f'reparacion:{a.id}'], 'forma_pago': 'efectivo', 'via_caja': ''})
        ap = AplicacionPagoComision.objects.get()
        self.assertIn(a.reparacion.codigo, ap.detalle_snapshot)
        self.assertNotIn('object (', ap.detalle_snapshot)

    def test_data_pendiente_sin_localizar(self):
        _crear_asignacion_reparacion(self.empleado, Decimal('1234.50'))
        PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('10.25'))
        resp = self.client.get(reverse('detalle_empleado', kwargs={'id': self.empleado.id}))
        html = resp.content.decode()
        self.assertIn('data-pendiente="1234.50"', html)
        self.assertIn('data-saldo-favor="10.25"', html)

    def test_valor_excel_seguro(self):
        for peligroso in ['=HYPERLINK("x")', '+1', '-1', '@SUM(A1)', '\tx', '\rx']:
            self.assertEqual(comisiones.valor_excel_seguro(peligroso), "'" + peligroso)
        self.assertEqual(comisiones.valor_excel_seguro('Maria'), 'Maria')
        self.assertEqual(comisiones.valor_excel_seguro(5), 5)

    def test_excel_no_exporta_formulas_de_texto_de_usuario(self):
        cli = make_cliente(nombres='=cmd', apellido_paterno='Lopez')
        _crear_asignacion_reparacion(self.empleado, Decimal('100'), cliente=cli)
        PagoComisionEmpleado.objects.create(
            empleado=self.empleado, monto=Decimal('5'), descripcion='=1+1')
        resp = self.client.get(reverse('exportar_devengaciones_empleado_excel', kwargs={'id': self.empleado.id}))
        wb = openpyxl.load_workbook(BytesIO(resp.content))
        valores = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
                   if isinstance(c.value, str)]
        self.assertFalse([v for v in valores if v.startswith('=')], valores)
        self.assertTrue(any(v.startswith("'=cmd") for v in valores))
        self.assertTrue(any(v.startswith("'=1+1") for v in valores))


class FifoRemanenteTests(TestCase):

    def test_monto_de_reparte_solo_el_remanente(self):
        pago = _PagoFake('100')
        aplicado = []
        acc = {'saldo': Decimal('80'), 'aplicar': lambda p, m: aplicado.append(m)}
        sobrante = comisiones.asignar_pagos_fifo([pago], [acc], monto_de=lambda p: Decimal('30'))
        self.assertEqual(aplicado, [Decimal('30')])
        self.assertEqual(sobrante, Decimal('0'))


class Backfill0066Tests(TestCase):
    """La función `poblar` de la 0066 corrida contra el registro real de
    modelos: snapshot legible e idempotencia por remanente."""

    def setUp(self):
        self.empleado = make_empleado()
        self.mig = importlib.import_module('misastreria.migrations.0066_backfill_aplicaciones_pago_comision')

    def _poblar(self):
        self.mig.poblar(django_apps, None)

    def test_reparte_fifo_con_snapshot_e_idempotente(self):
        a1 = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        a2 = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        pago = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('150'))
        self._poblar()
        self._poblar()  # segunda corrida: no duplica
        self.assertEqual(pago.aplicaciones.count(), 2)
        self.assertEqual(pago.aplicaciones.aggregate(s=Sum('monto'))['s'], Decimal('150'))
        ap1 = AplicacionPagoComision.objects.get(reparacion_empleado=a1)
        self.assertEqual(ap1.monto, Decimal('100'))
        self.assertIn(a1.reparacion.codigo, ap1.detalle_snapshot)
        self.assertNotIn('object (', ap1.detalle_snapshot)
        self.assertEqual(AplicacionPagoComision.objects.get(reparacion_empleado=a2).monto, Decimal('50'))

    def test_pago_parcialmente_aplicado_completa_solo_el_remanente(self):
        a1 = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        a2 = _crear_asignacion_reparacion(self.empleado, Decimal('100'))
        pago = PagoComisionEmpleado.objects.create(empleado=self.empleado, monto=Decimal('150'))
        AplicacionPagoComision.objects.create(pago=pago, reparacion_empleado=a1, monto=Decimal('100'))
        self._poblar()
        self.assertEqual(pago.aplicaciones.aggregate(s=Sum('monto'))['s'], Decimal('150'))
        self.assertEqual(AplicacionPagoComision.objects.get(reparacion_empleado=a2).monto, Decimal('50'))
